"""Export the Insights recommended projects (MetricRecommendation rows) with their
objective associations to an .xlsx file.

Each recommendation belongs to a metric (e.g. AOV, Close Rate), which maps to an
AEE element and the aligned Objective. Usage:
    python manage.py export_insights [--out path.xlsx] [--all]
"""
from django.core.management.base import BaseCommand

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# Metric objective-keyword -> AEE element (mirrors INSIGHT_METRICS['obj']).
KEYWORD_AEE = {
    'order value': 'expand_purchase',
    'shopper': 'attract_traffic',
    'close': 'engage_customers',
}
AEE_LABEL = {
    'attract_traffic': 'Attract Traffic',
    'engage_customers': 'Engage Customers',
    'expand_purchase': 'Expand Purchase',
}


class Command(BaseCommand):
    help = 'Export Insights recommendations + objective associations to an .xlsx file.'

    def add_arguments(self, parser):
        parser.add_argument('--out', default='insights_recommendations.xlsx',
                            help='Output .xlsx path (default: insights_recommendations.xlsx)')
        parser.add_argument('--all', action='store_true',
                            help='Include inactive recommendations too (default: active only).')

    def handle(self, *args, **opts):
        from app.models import MetricRecommendation
        from app.insights import _BY_KEY
        from strategy.models import Objective

        obj_by_aee = {o.aee_alignment: o for o in Objective.objects.exclude(aee_alignment='')}
        direction_label = {'win': 'Win', 'loss': 'Loss'}

        qs = MetricRecommendation.objects.all() if opts['all'] else \
            MetricRecommendation.objects.filter(active=True)
        qs = qs.order_by('metric', 'direction', 'order', 'id')

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Insights Recommendations'
        headers = ['Metric', 'Direction', 'Recommended action', 'AEE element', 'Associated objective']
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='4A1F8A')
            cell.alignment = Alignment(vertical='center')

        n = 0
        for r in qs:
            spec = _BY_KEY.get(r.metric, {})
            keyword = spec.get('obj', '')
            aee = KEYWORD_AEE.get(keyword, '')
            objective = obj_by_aee.get(aee)
            ws.append([
                spec.get('label', r.metric),
                direction_label.get(r.direction, r.direction),
                r.text,
                objective.get_aee_alignment_display() if objective else AEE_LABEL.get(aee, ''),
                objective.name if objective else '',
            ])
            n += 1

        for i, w in enumerate([26, 10, 62, 20, 30], start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = 'A2'
        wb.save(opts['out'])
        self.stdout.write(self.style.SUCCESS(f'Exported {n} recommendations to {opts["out"]}'))
